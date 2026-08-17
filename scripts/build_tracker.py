#!/usr/bin/env python3
"""
AI Engineering Lab, 24-week progress tracker generator.

Reads curriculum/manifest.json (the single source of truth for the program)
and emits curriculum/tracking/ai-engineering-lab-24-week-tracker.xlsx:

  * Dashboard, learner name, start date, per-week status + % bars, progress chart
  * Week 01..24, one sheet per week with day-by-day checklist, status dropdowns
                  (☐ Not started / ▶ In progress / ✅ Done / ⏭ Skipped),
                  conditional formatting, and an in-cell progress bar
  * About, how to use the workbook

Usage:
    python scripts/build_tracker.py
"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
MANIFEST_PATH = ROOT / "curriculum" / "manifest.json"
OUT_PATH = ROOT / "curriculum" / "tracking" / "ai-engineering-lab-24-week-tracker.xlsx"
LOGO_PATH = ROOT / "assets" / "zorost-logo.png"

# ---- Brand palette (Zorost) -------------------------------------------------
NAVY = "0F1B2D"
AMBER = "F5B301"
PAPER = "F7F8FA"
GREEN_FILL = "C6EFCE"
GREEN_FONT = "1E6B34"
AMBER_FILL = "FFF2CC"
AMBER_FONT = "8A6D00"
GREY_FILL = "E7E6E6"
GREY_FONT = "595959"
GRID = "D9DEE8"

STATUS_NOT = "☐ Not started"
STATUS_PROG = "▶ In progress"
STATUS_DONE = "✅ Done"
STATUS_SKIP = "⏭ Skipped"
STATUSES = [STATUS_NOT, STATUS_PROG, STATUS_DONE, STATUS_SKIP]

thin = Side(style="thin", color=GRID)
border = Border(left=thin, right=thin, top=thin, bottom=thin)

font_title = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
font_sub = Font(name="Calibri", size=11, italic=True, color="FFFFFF")
font_hdr = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
font_body = Font(name="Calibri", size=11, color="1A1A1A")
font_bold = Font(name="Calibri", size=11, bold=True, color="1A1A1A")
font_small = Font(name="Calibri", size=10, color="4A4A4A")

fill_navy = PatternFill("solid", fgColor=NAVY)
fill_amber = PatternFill("solid", fgColor=AMBER)
fill_paper = PatternFill("solid", fgColor=PAPER)

al_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
al_center = Alignment(horizontal="center", vertical="center", wrap_text=True)


def set_widths(ws, widths: dict[int, float]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w


def style_range(ws, cell_range: str, font=None, fill=None, border=None, alignment=None):
    for row in ws[cell_range]:
        for c in row:
            if font:
                c.font = font
            if fill:
                c.fill = fill
            if border:
                c.border = border
            if alignment:
                c.alignment = alignment


def build_week_sheet(wb: Workbook, week: dict, sheet_index: int) -> tuple[str, int]:
    """Build one week sheet; return (sheet name, summary row)."""
    name = f"Week {week['week']:02d}"
    ws = wb.create_sheet(name)

    # Header band
    ws.merge_cells("A1:F1")
    ws["A1"] = f"Week {week['week']:02d}, {week['title']}"
    ws["A1"].font = font_title
    ws["A1"].fill = fill_navy
    ws["A1"].alignment = al_left
    ws.row_dimensions[1].height = 26

    ws.merge_cells("A2:F2")
    ws["A2"] = f"Section: {week['section']} · Category: {week['category']} · ZoroLogistics case study"
    ws["A2"].font = font_sub
    ws["A2"].fill = fill_navy
    ws["A2"].alignment = al_left

    ws.merge_cells("A3:F3")
    ws["A3"] = f"🎯 Use case: {week['use_case']}"
    ws["A3"].font = font_bold
    ws["A3"].fill = fill_amber
    ws["A3"].alignment = al_left

    # Checklist header
    headers = ["Day", "Task", "Status", "Notes"]
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=5, column=col, value=h)
        c.font = font_hdr
        c.fill = fill_navy
        c.border = border
        c.alignment = al_center

    # Checklist rows
    first_data, last_data = 6, 5 + len(week["checklist"])
    for i, item in enumerate(week["checklist"]):
        r = first_data + i
        day_cell = ws.cell(row=r, column=1, value=item["day"])
        task_cell = ws.cell(row=r, column=2, value=item["task"])
        status_cell = ws.cell(row=r, column=3, value=STATUS_NOT if item["day"] != "Milestone" else STATUS_NOT)
        notes_cell = ws.cell(row=r, column=4, value="")
        for cell in (day_cell, task_cell, status_cell, notes_cell):
            cell.border = border
            cell.font = font_body
        day_cell.alignment = al_center
        day_cell.font = font_bold
        task_cell.alignment = al_left
        status_cell.alignment = al_center
        notes_cell.alignment = al_left
        if item["day"] == "Milestone":
            for cell in (day_cell, task_cell, status_cell, notes_cell):
                cell.fill = fill_paper
        ws.row_dimensions[r].height = 30

    # Status dropdown + conditional formatting
    dv = DataValidation(
        type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True,
        showDropDown=False, errorTitle="Invalid status",
        error="Pick one of: ☐ Not started, ▶ In progress, ✅ Done, ⏭ Skipped",
    )
    ws.add_data_validation(dv)
    dv.add(f"C{first_data}:C{last_data}")

    rng = f"C{first_data}:C{last_data}"
    ws.conditional_formatting.add(
        rng, FormulaRule(formula=[f'$C{first_data}="{STATUS_DONE}"'], fill=PatternFill("solid", fgColor=GREEN_FILL), font=Font(color=GREEN_FONT)))
    ws.conditional_formatting.add(
        rng, FormulaRule(formula=[f'$C{first_data}="{STATUS_PROG}"'], fill=PatternFill("solid", fgColor=AMBER_FILL), font=Font(color=AMBER_FONT)))
    ws.conditional_formatting.add(
        rng, FormulaRule(formula=[f'$C{first_data}="{STATUS_SKIP}"'], fill=PatternFill("solid", fgColor=GREY_FILL), font=Font(color=GREY_FONT)))

    # Completion summary
    sum_r = last_data + 1
    ws.cell(row=sum_r, column=1, value="Week completion").font = font_bold
    pct_cell = ws.cell(row=sum_r, column=2, value=f'=COUNTIF(C{first_data}:C{last_data},"{STATUS_DONE}")/{len(week["checklist"])}')
    pct_cell.number_format = "0%"
    pct_cell.font = font_bold
    pct_cell.border = border
    pct_cell.alignment = al_center
    bar_cell = ws.cell(
        row=sum_r, column=3,
        value=(f'=REPT("█",ROUND(B{sum_r}*20,0))&REPT("░",20-ROUND(B{sum_r}*20,0))'))
    bar_cell.alignment = al_left
    ws.merge_cells(start_row=sum_r, start_column=3, end_row=sum_r, end_column=4)
    ws.cell(row=sum_r, column=5, value="Deliverables: " + " · ".join(week["deliverables"])).font = font_small
    ws.merge_cells(start_row=sum_r, start_column=5, end_row=sum_r, end_column=6)

    set_widths(ws, {1: 12, 2: 78, 3: 18, 4: 34, 5: 40, 6: 40})
    ws.freeze_panes = "A6"
    ws.sheet_view.zoomScale = 100
    return name, sum_r


def build_dashboard(wb: Workbook, manifest: dict, week_sheets: list[tuple[str, int]], week_titles: list[tuple[int, str, str]]) -> None:
    ws = wb.active
    ws.title = "Dashboard"
    ws.sheet_properties.tabColor = AMBER

    if LOGO_PATH.exists():
        from openpyxl.drawing.image import Image as XLImage
        img = XLImage(str(LOGO_PATH))
        img.width = 80
        img.height = 80
        ws.add_image(img, "H2")

    ws.merge_cells("A1:F1")
    ws["A1"] = "AI Engineering Lab, 24-Week AI Engineering Program"
    ws["A1"].font = font_title
    ws["A1"].fill = fill_navy
    ws.row_dimensions[1].height = 26
    ws.merge_cells("A2:F2")
    ws["A2"] = f"{manifest['program']['tagline']} · by {manifest['program']['brand']} ({manifest['program']['brand_url']})"
    ws["A2"].font = font_sub
    ws["A2"].fill = fill_navy

    ws["B4"] = "Learner name"
    ws["B4"].font = font_bold
    ws["C4"].border = border
    ws["D4"] = "Start date"
    ws["D4"].font = font_bold
    ws["E4"].border = border
    ws["B5"] = "Program start"
    ws["B5"].font = font_small
    ws["C5"].border = border
    ws["D5"] = "Target end (24 weeks later)"
    ws["D5"].font = font_small
    ws["E5"].border = border

    # Table header
    hdr_row = 7
    for col, h in enumerate(["Week", "Section", "Title", "Status", "% Complete"], start=1):
        c = ws.cell(row=hdr_row, column=col, value=h)
        c.font = font_hdr
        c.fill = fill_navy
        c.border = border
        c.alignment = al_center

    first_data, last_data = hdr_row + 1, hdr_row + len(week_sheets)
    for i, ((wname, sum_row), (week_no, phase, title)) in enumerate(zip(week_sheets, week_titles)):
        r = first_data + i
        ws.cell(row=r, column=1, value=f"Week {week_no:02d}").font = font_body
        ws.cell(row=r, column=2, value=phase).font = font_body
        ws.cell(row=r, column=3, value=title).font = font_body
        status_cell = ws.cell(
            row=r, column=4,
            value=(f'=IF(\'{wname}\'!$B${sum_row}>=1,"✅ Complete",'
                   f'IF(\'{wname}\'!$B${sum_row}>0,"▶ In progress","☐ Not started"))'))
        status_cell.font = font_body
        pct_cell = ws.cell(row=r, column=5, value=f"='{wname}'!$B${sum_row}")
        pct_cell.number_format = "0%"
        pct_cell.font = font_body
        for col in range(1, 6):
            ws.cell(row=r, column=col).border = border
            ws.cell(row=r, column=col).alignment = al_left if col in (2, 3) else al_center

    pct_range = f"E{first_data}:E{last_data}"
    ws.conditional_formatting.add(
        pct_range,
        DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1,
                    color=AMBER, showValue=True))

    # Overall completion
    total_r = last_data + 2
    ws.cell(row=total_r, column=4, value="Overall completion").font = font_bold
    overall = ws.cell(row=total_r, column=5, value=f"=AVERAGE(E{first_data}:E{last_data})")
    overall.number_format = "0%"
    overall.font = Font(size=12, bold=True, color=GREEN_FONT)
    overall.border = border
    overall.alignment = al_center

    # Chart
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Completion by week"
    chart.y_axis.title = "% complete"
    chart.y_axis.numFmt = "0%"
    chart.y_axis.majorGridlines = None
    data = Reference(ws, min_col=5, min_row=hdr_row, max_row=last_data)
    cats = Reference(ws, min_col=1, min_row=first_data, max_row=last_data)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.width = 22
    chart.height = 10
    ws.add_chart(chart, f"G{hdr_row}")

    # Phase summary
    ph_r = total_r + 3
    ws.cell(row=ph_r, column=1, value="Phase summary").font = Font(size=13, bold=True, color="FFFFFF")
    ws.merge_cells(start_row=ph_r, start_column=1, end_row=ph_r, end_column=5)
    ws.cell(row=ph_r, column=1).fill = fill_navy
    for j, phase in enumerate(manifest["phases"]):
        r = ph_r + 1 + j
        first = first_data + sum(len(p["items"]) for p in manifest["phases"][:j])
        last = first + len(phase["items"]) - 1
        ws.cell(row=r, column=1, value=phase["name"]).font = font_bold
        ws.cell(row=r, column=2, value=f"Weeks {phase['weeks']}").font = font_small
        avg = ws.cell(row=r, column=3, value=f"=AVERAGE(E{first}:E{last})")
        avg.number_format = "0%"
        avg.font = font_body
        ws.cell(row=r, column=4, value=phase["description"]).font = font_small
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
        for col in range(1, 4):
            ws.cell(row=r, column=col).border = border

    # Legend
    lg_r = ph_r + len(manifest["phases"]) + 2
    ws.cell(row=lg_r, column=1, value="Legend").font = font_bold
    legend = [
        (STATUS_NOT, "not started yet"),
        (STATUS_PROG, "working on it"),
        (STATUS_DONE, "completed, pick this to tick it off ✅"),
        (STATUS_SKIP, "skipped deliberately (explain in Notes)"),
    ]
    for i, (status, meaning) in enumerate(legend):
        r = lg_r + 1 + i
        c = ws.cell(row=r, column=1, value=status)
        c.font = font_body
        c.alignment = al_center
        c.border = border
        fill_map = {STATUS_DONE: GREEN_FILL, STATUS_PROG: AMBER_FILL, STATUS_SKIP: GREY_FILL}
        if status in fill_map:
            c.fill = PatternFill("solid", fgColor=fill_map[status])
        ws.cell(row=r, column=2, value=meaning).font = font_small

    set_widths(ws, {1: 13, 2: 24, 3: 64, 4: 18, 5: 12, 6: 12, 7: 3, 8: 3})
    ws.freeze_panes = "A8"


def build_about(wb: Workbook, manifest: dict) -> None:
    ws = wb.create_sheet("About")
    ws.sheet_properties.tabColor = NAVY
    lines = [
        ("AI Engineering Lab, Progress Tracker", font_title, fill_navy),
        ("How to use this workbook", font_hdr, fill_navy),
        ("", font_body, None),
        ("1. Fill in your name and start date on the Dashboard sheet.", font_body, None),
        ("2. Each week has its own sheet (Week 01 … Week 24) with the week's checklist:", font_body, None),
        (" study, notebook, and use-case tasks, one line per program day.", font_body, None),
        ("3. Set a task's Status cell (dropdown) to ✅ Done when you finish it. The week's", font_body, None),
        (" completion % and the in-cell progress bar update automatically.", font_body, None),
        ("4. The Dashboard aggregates every week: status, % bars, phase averages, and a", font_body, None),
        (" progress chart. Tick all six tasks in a week to mark it ✅ Complete.", font_body, None),
        ("", font_body, None),
        ("Statuses", font_hdr, fill_navy),
        (f"{STATUS_NOT} · {STATUS_PROG} · {STATUS_DONE} · {STATUS_SKIP}", font_body, None),
        ("Skipped tasks count against the week %, use Notes to say why.", font_small, None),
        ("", font_body, None),
        ("Works in: Microsoft Excel, Google Sheets (File → Import), LibreOffice Calc.", font_body, None),
        ("The workbook is generated from curriculum/manifest.json by", font_body, None),
        ("scripts/build_tracker.py. Edit the manifest, never the xlsx by hand.", font_body, None),
        ("", font_body, None),
        (f"© 2026 {manifest['program']['brand']} LLC · {manifest['program']['brand_url']}", font_small, None),
    ]
    for i, (text, font, fill) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = font
        if fill:
            c.fill = fill
    set_widths(ws, {1: 110})
    ws.freeze_panes = "A3"


def main() -> None:
    manifest = json.loads(MANIFEST_PATH.read_text())
    wb = Workbook()

    week_sheets: list[tuple[str, int]] = []
    week_titles: list[tuple[int, str, str]] = []
    for phase in manifest["phases"]:
        for item in phase["items"]:
            wname, sum_row = build_week_sheet(wb, item, len(week_sheets))
            week_sheets.append((wname, sum_row))
            week_titles.append((item["week"], phase["name"], item["title"]))

    build_dashboard(wb, manifest, week_sheets, week_titles)
    build_about(wb, manifest)

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_PATH)
    print(f"Wrote {OUT_PATH} ({len(week_sheets)} week sheets + Dashboard + About)")


if __name__ == "__main__":
    main()
